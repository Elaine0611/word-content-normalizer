## 1. 是处理前的“结构检查器”。它只读取 DOCX，不修改文件。
## 2. 输出 JSON，列出文档的分节、段落、表格、页眉页脚、正文视觉对象数量和图片哈希；
## 3. 按原文顺序记录段落与表格，并标出识别到的标题层级、目录域和疑似修订记录表。

from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document

from docx_common import *


def inspect(path):
    doc = Document(path)

    blocks = []

    for i, b in enumerate(iter_blocks(doc)):

        # --------------------------------------------------
        # 表格
        # --------------------------------------------------
        if hasattr(b, "rows"):

            blocks.append({
                "index": i,
                "type": "table",
                "rows": len(b.rows),
                "columns": len(b.columns),
                "revision_score": revision_score(b),
                "header": [
                    c.text
                    for c in b.rows[0].cells
                ] if b.rows else [],
            })

            continue

        # --------------------------------------------------
        # 普通段落 / 图片段落
        # --------------------------------------------------

        if not b.text.strip() and not has_visual(b):
            continue

        heading = detect_block_heading_paragraph(b)
        
        item = {
            "index": i,
            "type": "paragraph",
            "style": paragraph_style_name(b),
            "text": b.text.strip(),
            "has_visual": has_visual(b),
            "toc_field": toc_field(b),

            # 标题识别结果
            "is_heading": heading is not None,
            "heading_level": (
                heading["level"]
                if heading
                else None
            ),
            "heading_kind": (
                heading["kind"]
                if heading
                else None
            ),
            "heading_number": (
                heading["number"]
                if heading
                else None
            ),
        }

        blocks.append(item)

    return {
        "file": str(
            Path(path).resolve()
        ),
        "sections": len(doc.sections),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "body_visuals": body_visual_count(doc),
        "image_hashes": image_hashes(doc),
        "headers": [
            " | ".join(
                p.text
                for p in s.header.paragraphs
            )
            for s in doc.sections
        ],
        "footers": [
            " | ".join(
                p.text
                for p in s.footer.paragraphs
            )
            for s in doc.sections
        ],
        "blocks": blocks,
    }


def main():
    ap = argparse.ArgumentParser()

    ap.add_argument(
        "input",
        type=Path,
    )

    ap.add_argument(
        "--json",
        type=Path,
    )

    a = ap.parse_args()

    result = inspect(a.input)

    payload = json.dumps(
        result,
        ensure_ascii=False,
        indent=2,
    )

    if a.json:
        a.json.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        a.json.write_text(
            payload,
            encoding="utf-8",
        )

    print(payload)


if __name__ == "__main__":
    main()
