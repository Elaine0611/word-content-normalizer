from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from docx import Document
from lxml import etree

from docx_common import (
    body_visual_count,
    image_hashes,
    iter_blocks,
    norm,
    paragraph_style_name,
    revision_score,
    revision_table,
    is_revision_heading,
    toc_field,
)
from segment_docx import verify_in_document_segments


def verify(src, dst):
    """Verify a complete normalized copy against its source document."""

    source = Document(src)
    output = Document(dst)
    toc = [
        p.text
        for p in output.paragraphs
        if (
            norm(p.text) == "目录"
            or paragraph_style_name(p).lower().startswith("toc")
            or toc_field(p)
        )
    ]
    revision_headings = [p.text for p in output.paragraphs if is_revision_heading(p.text)]
    revisions = []
    blocks = list(iter_blocks(output))
    for index, block in enumerate(blocks):
        if not hasattr(block, "rows"):
            continue
        previous = next(
            (item.text for item in reversed(blocks[:index])
             if hasattr(item, "text") and item.text.strip()),
            "",
        )
        if revision_table(block, previous):
            revisions.append(revision_score(block))
    report = {
        "source": str(src.resolve()),
        "output": str(dst.resolve()),
        "source_body_visuals": body_visual_count(source),
        "output_body_visuals": body_visual_count(output),
        "body_visuals_preserved": (
            body_visual_count(source) == body_visual_count(output)
        ),
        "media_bytes_preserved": (
            image_hashes(source) == image_hashes(output)
        ),
        "headers_empty": all(
            not "".join(p.text for p in section.header.paragraphs).strip()
            for section in output.sections
        ),
        "footers_empty": all(
            not "".join(p.text for p in section.footer.paragraphs).strip()
            for section in output.sections
        ),
        "toc_removed": not toc,
        "revision_tables_removed": not revisions,
        "revision_headings_removed": not revision_headings,
        "remaining_toc": toc,
        "remaining_revision_scores": revisions,
        "remaining_revision_headings": revision_headings,
    }
    report.update(verify_in_document_segments(output))
    report["passed"] = all(
        report[key]
        for key in (
            "body_visuals_preserved",
            "media_bytes_preserved",
            "headers_empty",
            "footers_empty",
            "toc_removed",
            "revision_tables_removed",
            "revision_headings_removed",
            "segment_coverage_passed",
            "segment_boundaries_marked",
        )
    )
    return report


def _block_hash(block):
    payload = etree.tostring(
        block._element,
        encoding="utf-8",
        with_tail=False,
    )
    return hashlib.sha256(payload).hexdigest()


def _block_visual_count(block):
    xml = block._element.xml
    return sum(
        xml.count(tag)
        for tag in ("<w:drawing", "<w:pict", "<w:object")
    )


def verify_segment(src, dst, expected_indices):
    """Verify that an exported DOCX contains exactly the selected blocks."""

    source = Document(src)
    output = Document(dst)
    source_blocks = list(iter_blocks(source))
    output_blocks = list(iter_blocks(output))
    expected_indices = list(expected_indices)
    indices_valid = (
        expected_indices == sorted(set(expected_indices))
        and all(0 <= index < len(source_blocks) for index in expected_indices)
    )

    if indices_valid:
        expected_blocks = [source_blocks[index] for index in expected_indices]
        expected_hashes = [_block_hash(block) for block in expected_blocks]
        expected_visuals = sum(
            _block_visual_count(block) for block in expected_blocks
        )
    else:
        expected_hashes = []
        expected_visuals = None

    output_hashes = [_block_hash(block) for block in output_blocks]
    report = {
        "source": str(src.resolve()),
        "output": str(dst.resolve()),
        "expected_indices": expected_indices,
        "indices_valid": indices_valid,
        "expected_block_count": len(expected_indices),
        "output_block_count": len(output_blocks),
        "blocks_match": indices_valid and expected_hashes == output_hashes,
        "expected_body_visuals": expected_visuals,
        "output_body_visuals": body_visual_count(output),
        "body_visuals_match": (
            indices_valid
            and expected_visuals == body_visual_count(output)
        ),
        "media_bytes_preserved": (
            image_hashes(source) == image_hashes(output)
        ),
    }
    report["passed"] = all(
        report[key]
        for key in (
            "indices_valid",
            "blocks_match",
            "body_visuals_match",
            "media_bytes_preserved",
        )
    )
    return report


def _parse_indices(value):
    if not value.strip():
        return []
    return [int(item.strip()) for item in value.split(",")]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--json", type=Path)
    parser.add_argument(
        "--segment-indices",
        help="Comma-separated source body block indices for segment mode.",
    )
    args = parser.parse_args()

    if args.segment_indices is None:
        report = verify(args.source, args.output)
    else:
        report = verify_segment(
            args.source,
            args.output,
            _parse_indices(args.segment_indices),
        )

    payload = json.dumps(report, ensure_ascii=False, indent=2)

    if args.json:
        args.json.parent.mkdir(parents=True, exist_ok=True)
        args.json.write_text(payload, encoding="utf-8")

    print(payload)
    raise SystemExit(0 if report["passed"] else 1)


if __name__ == "__main__":
    main()
