## 实际清理流程
##负责复制文档、清空页眉页脚、整理封面、删除修订表和目录、调整标题样式、删除空段落，最后保存新的 DOCX。

from __future__ import annotations
import argparse,re,shutil
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.text.paragraph import Paragraph
from docx.shared import Pt
from docx_common import *
from segment_docx import apply_segments_in_place

DEFS={norm(x) for x in ("名称解释","术语解释","名词解释","定义")}

def definition_heading(text):
    value=norm(text).rstrip(":：")
    return value in DEFS

def heading_level(p):
    info=detect_heading_paragraph(p)
    if not info:return None
    level=info.get("level")
    return level if isinstance(level,int) and level>0 else None

def toc_style(p):
    return paragraph_style_name(p).lower().startswith("toc")

def standard_heading_style(doc,level):
    for name in (f"Heading {level}",f"标题 {level}"):
        try:return doc.styles[name]
        except KeyError:pass
    return None

def remove(el):
    if el.getparent() is not None:el.getparent().remove(el)

def clear_part(part):
    for child in list(part._element):remove(child)
    part.add_paragraph("")

def replace_text(p,text):
    if has_visual(p):raise RuntimeError("Refusing to rewrite a visual anchor paragraph")
    for run in list(p.runs):remove(run._element)
    p.add_run(text)

def flatten_cover(doc):
    ps=list(doc.paragraphs);stop=None
    for i,p in enumerate(ps):
        if definition_heading(p.text) or norm(p.text)=="目录" or heading_level(p)==1:stop=i;break
    if not stop:return
    cover=ps[:stop]
    if any(has_visual(p) for p in cover):return
    texts=[p.text.strip() for p in cover if p.text.strip()];out=[];chars=[]
    for text in texts:
        if is_revision_heading(text):continue
        if len(text)==1 and text.isalpha():chars.append(text);continue
        if chars:out.append("".join(chars));chars=[]
        value=re.sub(r"版本\s*[-—_]\s*(\d+)",r"版本\1",text.strip("【】[] "),flags=re.I)
        nv=norm(value)
        if nv and any(nv in norm(x) or norm(x) in nv for x in out):continue
        out.append(value)
    if chars:out.append("".join(chars))
    out=[x for x in out if x]
    if not out:return
    replace_text(cover[0]," ".join(out));cover[0].style=doc.styles["Normal"];cover[0].alignment=WD_ALIGN_PARAGRAPH.LEFT;cover[0].paragraph_format.space_after=Pt(12)
    for p in cover[1:]:remove(p._p)

def remove_revisions(doc):
    for block in list(iter_blocks(doc)):
        if not hasattr(block,"rows"):continue
        prev=block._element.getprevious();nearby=[]
        while prev is not None and prev.tag.endswith("}p"):
            text=Paragraph(prev,doc).text.strip()
            nearby.append((prev,text))
            if text:break
            prev=prev.getprevious()
        title=nearby[-1][1] if nearby else ""
        if not revision_table(block,title):continue
        remove(block._element)
        if is_revision_heading(title):
            for element,_ in nearby:remove(element)

def remove_toc(doc):
    ps=list(doc.paragraphs);start=None
    for i,p in enumerate(ps):
        if norm(p.text)=="目录" or toc_field(p):start=i;break
    if start is None:return
    for p in ps[start:]:
        is_toc_style=toc_style(p)
        if p is not ps[start] and heading_level(p)==1 and not is_toc_style:break
        if has_visual(p):raise RuntimeError("TOC boundary contains a visual; manual review required")
        if p is ps[start] or is_toc_style or toc_field(p) or not p.text.strip():remove(p._p)

def normalize_defs(doc,spacing):
    active=False
    for p in list(doc.paragraphs):
        n=norm(p.text)
        if definition_heading(p.text):active=True;p.style=doc.styles["Normal"];p.paragraph_format.space_after=Pt(spacing);continue
        if active and (n=="目录" or heading_level(p)==1):active=False
        if active and p.text.strip():
            if p._p.pPr is not None and p._p.pPr.numPr is not None:p._p.pPr.remove(p._p.pPr.numPr)
            p.style=doc.styles["Normal"];p.paragraph_format.left_indent=Pt(0);p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.space_after=Pt(spacing)

def bold_item(p):
    runs=[r for r in p.runs if r.text.strip()]
    return bool(p.text.strip() and len(p.text.strip())<=40 and runs and all(r.bold for r in runs))

def normalize_hierarchy(doc,spacing):
    ps=list(doc.paragraphs)
    for i,p in enumerate(ps):
        if not p.text.strip() or has_visual(p):continue
        level=heading_level(p)
        if level:
            style=standard_heading_style(doc,level)
            if style is not None:p.style=style
            p.paragraph_format.space_before=Pt(spacing);p.paragraph_format.space_after=Pt(spacing)
        elif bold_item(p):
            p.paragraph_format.space_before=Pt(spacing);p.paragraph_format.space_after=Pt(0)
            if i+1<len(ps) and ps[i+1].text.strip() and not has_visual(ps[i+1]):ps[i+1].paragraph_format.space_after=Pt(spacing)

def remove_safe_blanks(doc):
    protected_tags=(
        "<w:br","<w:sectPr","<w:fldSimple","<w:instrText","<w:bookmarkStart","<w:bookmarkEnd",
        "<w:commentRangeStart","<w:commentRangeEnd","<w:commentReference","<w:footnoteReference",
        "<w:endnoteReference","<w:sdt","<w:ins","<w:del",
    )

    def removable_blank(p):
        if p.text.strip() or has_visual(p):return False
        return not any(tag in p._p.xml for tag in protected_tags)

    for p in list(doc.paragraphs):
        if not removable_blank(p):continue
        before=p._p.getprevious();after=p._p.getnext()
        if before is not None and after is not None and before.tag.endswith("}p") and after.tag.endswith("}p"):remove(p._p)

    while doc.paragraphs:
        p=doc.paragraphs[-1]
        if not removable_blank(p):break
        before=p._p.getprevious()
        if before is None or not before.tag.endswith("}p"):break
        remove(p._p)

def clean(src,dst,spacing):
    if src.resolve()==dst.resolve():raise ValueError("Source overwrite is disabled")
    if dst.exists():raise FileExistsError(f"Output already exists: {dst}")
    doc=Document(src)
    for s in doc.sections:clear_part(s.header);clear_part(s.footer)
    remove_revisions(doc);flatten_cover(doc);normalize_defs(doc,spacing);remove_toc(doc);normalize_hierarchy(doc,spacing);remove_safe_blanks(doc)
    apply_segments_in_place(doc)
    dst.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,dst);doc.save(dst)

def main():
    ap=argparse.ArgumentParser();ap.add_argument("input",type=Path);ap.add_argument("--spacing-pt",type=float,default=12);a=ap.parse_args();dst=default_output(a.input);clean(a.input,dst,a.spacing_pt);print(dst.resolve())
if __name__=="__main__":main()
