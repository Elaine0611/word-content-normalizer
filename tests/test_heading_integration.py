import sys
import tempfile
import unittest
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from clean_docx import clean
from docx_common import default_output, detect_heading
from verify_docx import verify


class HeadingIntegrationTests(unittest.TestCase):
    def test_default_output_uses_source_name_date_and_next_version(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "TCR系统使用说明.docx"
            today = date(2026, 10, 25)
            first = default_output(source, today)
            self.assertEqual(first.name, "TCR系统使用说明_规范化_20261025-V1.docx")
            first.touch()
            self.assertEqual(default_output(source, today).name,
                             "TCR系统使用说明_规范化_20261025-V2.docx")
            (source.parent / "TCR系统使用说明_规范化_20261025-V3.docx").touch()
            self.assertEqual(default_output(source, today).name,
                             "TCR系统使用说明_规范化_20261025-V4.docx")
            self.assertEqual(default_output(source, date(2026, 10, 26)).name,
                             "TCR系统使用说明_规范化_20261026-V1.docx")

    def test_existing_explicit_output_is_not_overwritten(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.docx"
            output = Path(temp_dir) / "existing.docx"
            Document().save(source)
            output.write_bytes(b"keep existing output")
            with self.assertRaises(FileExistsError):
                clean(source, output, 12)
            self.assertEqual(output.read_bytes(), b"keep existing output")

    def clean_document(self, build):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        source = Path(temp_dir.name) / "source.docx"
        output = Path(temp_dir.name) / "output.docx"
        doc = Document()
        build(doc)
        doc.save(source)
        clean(source, output, 12)
        return source, output, Document(output)

    def test_custom_style_inheriting_heading_three_is_normalized(self):
        def build(doc):
            custom = doc.styles.add_style(
                "Product Subheading",
                WD_STYLE_TYPE.PARAGRAPH,
            )
            custom.base_style = doc.styles["Heading 3"]
            doc.add_paragraph("Configuration", style=custom)

        _, _, output = self.clean_document(build)

        self.assertEqual(output.paragraphs[1].style.name, "Heading 3")

    def test_decimal_heading_level_three_uses_heading_three(self):
        def build(doc):
            doc.add_paragraph("1.2.3 Configuration")

        _, _, output = self.clean_document(build)

        self.assertEqual(output.paragraphs[1].style.name, "Heading 3")

    def test_year_prefix_is_not_treated_as_heading(self):
        self.assertIsNone(detect_heading("2024年生产报告"))
        self.assertEqual(detect_heading("1. 目的")["level"], 1)
        self.assertEqual(detect_heading("1.1配置")["level"], 2)

    def test_word_heading_one_ends_toc_removal(self):
        def build(doc):
            toc_style = doc.styles.add_style("TOC 1", WD_STYLE_TYPE.PARAGRAPH)
            doc.add_paragraph("目录")
            doc.add_paragraph("1. Entry", style=toc_style)
            doc.add_paragraph("Overview", style="Heading 1")
            doc.add_paragraph("Body")
            doc.add_paragraph("Keep this styled paragraph", style=toc_style)

        _, _, output = self.clean_document(build)
        texts = [p.text for p in output.paragraphs]

        self.assertNotIn("目录", texts)
        self.assertNotIn("1. Entry", texts)
        self.assertIn("##", texts)
        self.assertIn("Overview", texts)
        self.assertIn("Keep this styled paragraph", texts)

    def test_verify_accepts_unresolved_paragraph_style(self):
        def build(doc):
            paragraph = doc.add_paragraph("Body")
            paragraph._p.get_or_add_pPr().get_or_add_pStyle().val = "MissingStyle"

        source, output, _ = self.clean_document(build)

        self.assertTrue(verify(source, output)["passed"])

    def test_long_section_is_split_inside_one_verified_docx(self):
        def build(doc):
            doc.add_paragraph("一、数据种类")
            doc.add_paragraph("甲" * 600)
            doc.add_paragraph("乙" * 300)

        source, output, result = self.clean_document(build)
        report = verify(source, output)

        self.assertTrue(report["passed"])
        self.assertEqual(report["segment_count"], 2)
        self.assertEqual([p.text for p in result.paragraphs], [
            "##", "一、数据种类", "甲" * 600, "##", "乙" * 300,
        ])
        self.assertFalse(bool(result.paragraphs[3].paragraph_format.page_break_before))

    def test_multiple_sections_are_marked_in_one_verified_docx(self):
        def build(doc):
            doc.add_paragraph("一、第一段")
            doc.add_paragraph("甲" * 200)
            doc.add_section()
            doc.add_paragraph("二、第二段")
            doc.add_paragraph("乙" * 200)

        source, output, result = self.clean_document(build)
        report = verify(source, output)

        self.assertEqual(len(Document(source).sections), 2)
        self.assertTrue(report["passed"])
        self.assertEqual(report["segment_count"], 2)
        self.assertEqual([p.text for p in result.paragraphs if p.text == "##"],
                         ["##", "##"])

    def test_short_heading_blocks_share_one_document_segment(self):
        def build(doc):
            doc.add_paragraph("一、First")
            doc.add_paragraph("甲" * 200)
            doc.add_paragraph("二、Second")
            doc.add_paragraph("乙" * 30)
            doc.add_paragraph("三、Third")
            doc.add_paragraph("丙" * 120)

        source, output, result = self.clean_document(build)
        report = verify(source, output)

        self.assertTrue(report["passed"])
        self.assertEqual(report["segment_count"], 2)
        self.assertEqual([p.text for p in result.paragraphs if p.text == "##"], ["##", "##"])
        self.assertEqual(result.paragraphs[0].text, "##")
        self.assertEqual(result.paragraphs[3].text, "##")
        self.assertFalse(bool(result.paragraphs[3].paragraph_format.page_break_before))

    def test_short_block_does_not_end_with_next_title_only(self):
        def build(doc):
            doc.add_paragraph("一、开始")
            doc.add_paragraph("甲" * 144)
            doc.add_paragraph("二、流程")
            doc.add_paragraph("三、筛选")
            doc.add_paragraph("乙" * 20)
            doc.add_paragraph("四、结束")
            doc.add_paragraph("丙" * 200)

        source, output, result = self.clean_document(build)
        report = verify(source, output)
        texts = [p.text for p in result.paragraphs]

        self.assertTrue(report["passed"])
        self.assertEqual(report["segment_count"], 2)
        self.assertEqual(report["segment_char_counts"][0], 176)
        self.assertEqual(texts, [
            "##", "一、开始", "甲" * 144, "二、流程", "三、筛选",
            "乙" * 20, "##", "四、结束", "丙" * 200,
        ])

    def test_table_start_is_marked_without_splitting_table(self):
        def build(doc):
            doc.add_paragraph("一、数据种类")
            doc.add_paragraph("甲" * 600)
            table = doc.add_table(rows=1, cols=1)
            table.cell(0, 0).text = "乙" * 300

        source, output, result = self.clean_document(build)
        report = verify(source, output)

        self.assertTrue(report["passed"])
        self.assertEqual(report["segment_count"], 2)
        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.paragraphs[-1].text, "##")
        self.assertEqual(result.tables[0].cell(0, 0).text, "乙" * 300)

    def test_heading_five_steps_stay_with_primary_heading_when_merging(self):
        def build(doc):
            doc.add_paragraph("封" * 110)
            doc.add_paragraph("一、筛选TCR")
            doc.add_paragraph("甲" * 20)
            doc.add_paragraph("Site", style="Heading 5")
            doc.add_paragraph("乙" * 20)
            doc.add_paragraph("Business", style="Heading 5")
            doc.add_paragraph("点击 Search 进行筛选")
            doc.add_paragraph("二、列表介绍")
            doc.add_paragraph("丙" * 200)

        source, output, result = self.clean_document(build)
        texts = [p.text for p in result.paragraphs]
        self.assertTrue(verify(source, output)["passed"])
        self.assertEqual(texts.count("##"), 2)
        self.assertEqual(texts[texts.index("##", 1) - 1], "点击 Search 进行筛选")
        self.assertEqual(texts[texts.index("##", 1) + 1], "二、列表介绍")

    def test_long_block_prefers_minor_heading_to_plain_paragraph_edge(self):
        def build(doc):
            doc.add_paragraph("一、流程")
            doc.add_paragraph("甲" * 600)
            doc.add_paragraph("Site", style="Heading 5")
            doc.add_paragraph("乙" * 300)

        source, output, result = self.clean_document(build)
        texts = [p.text for p in result.paragraphs]
        self.assertTrue(verify(source, output)["passed"])
        self.assertEqual(texts, ["##", "一、流程", "甲" * 600,
                                 "##", "Site", "乙" * 300])

    def test_long_plain_paragraph_splits_at_sentence_end(self):
        def build(doc):
            doc.add_paragraph("一、说明")
            paragraph = doc.add_paragraph()
            paragraph.add_run("甲" * 450 + "。").bold = True
            paragraph.add_run("乙" * 450 + "。")

        source, output, result = self.clean_document(build)
        texts = [p.text for p in result.paragraphs]
        self.assertTrue(verify(source, output)["passed"])
        self.assertEqual(texts, ["##", "一、说明", "甲" * 450 + "。",
                                 "##", "乙" * 450 + "。"])
        self.assertTrue(result.paragraphs[2].runs[0].bold)

    def test_english_revision_history_is_removed_before_cover_cleanup(self):
        def build(doc):
            doc.add_paragraph("TCR系统开发使用说明")
            doc.add_paragraph("Revision History")
            table = doc.add_table(rows=2, cols=4)
            for cell, value in zip(table.rows[0].cells,
                                   ["Date", "Version No.", "Comments", "Author"]):
                cell.text = value
            for cell, value in zip(table.rows[1].cells,
                                   ["2023-05-19", "1.0", "Initial Release", "Editor"]):
                cell.text = value
            doc.add_paragraph("1.0目的")
            doc.add_paragraph("正文内容")

        source, output, result = self.clean_document(build)
        texts = [p.text for p in result.paragraphs]
        self.assertEqual(len(result.tables), 0)
        self.assertEqual(texts[:2], ["##", "TCR系统开发使用说明"])
        self.assertNotIn("Revision History", " ".join(texts))
        self.assertTrue(verify(source, output)["passed"])

    def test_chinese_revision_table_removed_and_preamble_counted_before_merge(self):
        def build(doc):
            doc.add_paragraph("封" * 50)
            doc.add_paragraph("版本变更")
            table = doc.add_table(rows=2, cols=4)
            for cell, value in zip(table.rows[0].cells,
                                   ["版本", "变更内容", "修订人", "修订时间"]):
                cell.text = value
            for cell, value in zip(table.rows[1].cells,
                                   ["v1.0", "初版", "编辑", "2023/5/15"]):
                cell.text = value
            doc.add_paragraph("一、甲")
            doc.add_paragraph("甲" * 100)
            doc.add_paragraph("二、乙")
            doc.add_paragraph("乙" * 200)

        source, output, result = self.clean_document(build)
        report = verify(source, output)

        self.assertTrue(report["passed"])
        self.assertEqual(len(result.tables), 0)
        self.assertNotIn("版本变更", [p.text for p in result.paragraphs])
        self.assertEqual(report["segment_count"], 2)
        self.assertEqual(report["segment_char_counts"][0], 153)
        self.assertEqual([p.text for p in result.paragraphs if p.text == "##"],
                         ["##", "##"])

    def test_business_table_with_date_and_version_is_preserved(self):
        def build(doc):
            doc.add_paragraph("一、设备数据")
            table = doc.add_table(rows=1, cols=3)
            for cell, value in zip(table.rows[0].cells,
                                   ["设备版本", "发布日期", "设备名称"]):
                cell.text = value

        source, output, result = self.clean_document(build)
        self.assertEqual(len(result.tables), 1)
        self.assertEqual(result.tables[0].rows[0].cells[2].text, "设备名称")
        self.assertTrue(verify(source, output)["passed"])


if __name__ == "__main__":
    unittest.main()
