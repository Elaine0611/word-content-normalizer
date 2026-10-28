import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document


SCRIPTS = Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))

from export_segments import export_segments
from segment_docx import count_chars, segment_for_document, segment_for_export


class SegmentExportTests(unittest.TestCase):
    def test_word_count_keeps_ninety_six_word_tcr_passage_short(self):
        doc = Document()
        doc.add_paragraph("四、其他", style="Heading 1")
        doc.add_paragraph(
            "如下所示，对于 Status 为 Release 的TCR（筛选 Introduce Status → "
            "Need My Verification），则是需要 TE 进行验证，根据验证的结果再确定是否导入"
        )
        doc.add_paragraph("导入后记录", style="Heading 1")
        doc.add_paragraph(
            "在Teams LSSC_DT_TE 共享文件LSSC DT TE Test Code update list，"
            "是TCR 导入的记录表，方便后续的反查和追溯，如下所示："
        )
        doc.add_paragraph("Sheet", style="Heading 1")
        doc.add_paragraph("甲" * 200)

        plan = segment_for_document(doc)

        self.assertEqual(count_chars(" ".join(
            item["text"]
            for block in plan["initial_blocks"][:2]
            for item in block["elements"]
        )), 96)
        self.assertEqual(plan["merged_block_count"], 1)
        self.assertEqual(plan["merged_blocks"][0]["source_block_indices"], [0, 1, 2])

    def make_paths(self):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        root = Path(temp_dir.name)
        return root / "source.docx", root / "segments"

    def test_exports_contiguous_blocks_and_preserves_table(self):
        source, output_dir = self.make_paths()
        doc = Document()
        doc.add_paragraph("Document preamble")
        doc.add_paragraph("一、First")
        doc.add_paragraph("甲" * 200)
        doc.add_paragraph("二、Second")
        doc.add_paragraph("乙" * 200)
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Key"
        table.cell(0, 1).text = "Value"
        doc.save(source)

        report = export_segments(source, output_dir)

        self.assertTrue(report["passed"])
        self.assertTrue(report["coverage_passed"])
        self.assertEqual(report["segment_count"], 2)
        self.assertEqual(report["covered_source_indices"], list(range(6)))

        first = Document(report["segments"][0]["output"])
        second = Document(report["segments"][1]["output"])
        first_text = [paragraph.text for paragraph in first.paragraphs]
        second_text = [paragraph.text for paragraph in second.paragraphs]

        self.assertIn("Document preamble", first_text)
        self.assertIn("一、First", first_text)
        self.assertNotIn("二、Second", first_text)
        self.assertIn("二、Second", second_text)
        self.assertEqual(len(second.tables), 1)
        self.assertEqual(second.tables[0].cell(0, 1).text, "Value")

    def test_long_content_splits_only_between_paragraphs(self):
        source, _ = self.make_paths()
        doc = Document()
        doc.add_paragraph("一、Long section")
        doc.add_paragraph("甲" * 600)
        doc.add_paragraph("乙" * 300)
        doc.save(source)

        plan = segment_for_export(source)
        segments = plan["final_blocks"]
        all_indices = [
            item["index"]
            for segment in segments
            for item in segment["elements"]
        ]

        self.assertEqual(len(segments), 2)
        self.assertEqual(all_indices, [0, 1, 2])
        self.assertTrue(
            all(
                "split_fragment" not in item
                for segment in segments
                for item in segment["elements"]
            )
        )

    def test_single_long_paragraph_stays_with_its_heading(self):
        source, _ = self.make_paths()
        doc = Document()
        doc.add_paragraph("一、Long section")
        doc.add_paragraph("甲" * 1200)
        doc.save(source)

        plan = segment_for_export(source)

        self.assertEqual(plan["final_block_count"], 1)
        self.assertGreater(plan["final_blocks"][0]["char_count"], 800)

    def test_short_block_can_merge_up_to_one_thousand_characters(self):
        source, _ = self.make_paths()
        doc = Document()
        doc.add_paragraph("一、Short")
        doc.add_paragraph("甲" * 100)
        doc.add_paragraph("二、Long")
        doc.add_paragraph("乙" * 850)
        doc.save(source)

        plan = segment_for_export(source)

        self.assertEqual(plan["initial_block_count"], 2)
        self.assertEqual(plan["merged_block_count"], 1)
        self.assertGreater(plan["merged_blocks"][0]["char_count"], 800)
        self.assertLessEqual(plan["merged_blocks"][0]["char_count"], 1000)
        self.assertEqual(plan["final_block_count"], 2)

    def test_multiple_sections_are_refused(self):
        source, _ = self.make_paths()
        doc = Document()
        doc.add_paragraph("一、First")
        doc.add_section()
        doc.add_paragraph("二、Second")
        doc.save(source)

        with self.assertRaisesRegex(ValueError, "single-section"):
            segment_for_export(source)


if __name__ == "__main__":
    unittest.main()
